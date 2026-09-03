# SPDX-License-Identifier: GPL-2.0-only

"""Script to parse the IO sheet and generate the linux_drivers/pinctrl/google/pinctrl-google-mbu.c file.

The sheet needs to be downloaded from go/<target>-io and should be saved in the
same directory as that of the parser.

Instructions how to download sswrps from si_codex to generate missing
pins registers can be found in go/mbu_parser_sswrps

Current sheet date: July 9, 2024.

Pre-requisite:
Excel parser module. 'pip install openpyxl' if not exists.

Arguments: [OPTIONAL] --out_file: Writes the output to the filename provided
"""

import argparse
import collections
import json
import sys
import textwrap

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

_MAX_LINE_LENGTH = 100
_TAB_SIZE = 8


class Pin:
  """Represents a physical pin/pad."""

  def __init__(self, pad_nr, pad_name, pin_functions):
    self.pad_nr = int(pad_nr.split("_")[1])
    self.pad_name = pad_name
    self.pin_functions = pin_functions


class TextWrapper:
  r"""Custom text wrapper for handling tabs and line lengths.

  This class extends the functionality of the standard textwrap.TextWrapper
  class to provide more control over tab sizes and indentation. It allows
  for specifying a custom tab size and ensures that tabs are preserved
  in the output, even when wrapping text. Without it, textwrap.TextWrapper
  will count `\t` as 1 character, which may not be desirable as it should be
  equal to 8 spaces according to kernel style guide.

  Attributes:
    tab_size: An integer representing the number of spaces for a single tab.
    tab: A string of spaces representing a single tab, calculated based on the
      tab_size.
    wrapper: An instance of textwrap.TextWrapper with customized indentation
      settings, using the calculated tab string.

  The class provides a `wrap` method that takes a string as input and returns
  a list of strings, where each string represents a wrapped line of text.
  The method utilizes the customized textwrap.TextWrapper instance to perform
  the wrapping while preserving tabs in the output.
  """

  def __init__(self, **kwargs):
    self.tab_size = kwargs.get("tab_size", _TAB_SIZE)
    self.tab = " " * self.tab_size

    kwargs["initial_indent"] = kwargs.get("initial_indent", "").replace(
        "\t", self.tab
    )
    kwargs["subsequent_indent"] = kwargs.get("subsequent_indent", "").replace(
        "\t", self.tab
    )

    self.wrapper = textwrap.TextWrapper(**kwargs)

  def wrap(self, text):
    """Wrap text while preserving tabs.

    This method wraps the input text into multiple lines, ensuring that
    the specified tab size is maintained and tabs are preserved in the
    output. It utilizes the customized textwrap.TextWrapper instance
    to perform the wrapping.

    Args:
      text: The string of text to be wrapped.

    Returns:
      A list of strings, where each string represents a wrapped line of
      text with tabs preserved.
    """
    text = self.wrapper.wrap(text)

    text = [t.replace(self.tab, "\t") for t in text]

    return text


def custom_container_get(container, el):
  if isinstance(container, dict):
    return container[el]
  elif isinstance(container, list):
    for l in container:
      if isinstance(l, dict) and l["name"] == el:
        return l
    return None
  else:
    raise TypeError(f"Unexpected type: {container}")


class Sswrp:

  """Represents a subsystem wrapper(sswrp)."""
  _SI_CODEX_FILENAME = {
      "aoss": "sswrp_aoss.json",
      "aoss_audio": "sswrp_aoss.json",
      "gdmc": "sswrp_gdmc.json",
      "hsios": "sswrp_hsio_s.json",
      "hsios_stby": "sswrp_hsio_s.json",
      "lsioe": "sswrp_lsioe.json",
      "lsios": "sswrp_lsios.json",
      "pcie": "sswrp_pcie.json",
  }

  _SI_CODEX_PATH = {
      "aoss": ["chip", "chip", "SSWRP_AOSS_AON", "block", "AOSS_BK_csr"],
      "aoss_audio": [
          "chip",
          "chip",
          "SSWRP_AOSS_AON",
          "block",
          "AOSS_AUDIO_BK_csr",
      ],
      "gdmc": ["chip", "block", "gpio_csrs"],
      "hsios": ["chip", "block", "HSIOS_BK_csr"],
      "hsios_stby": ["chip", "block", "HSIOS_STBY_BK_csr"],
      "lsioe": ["chip", "block", "LSIOE_BK_csr"],
      "lsios": ["chip", "block", "LSIOS_BK_csr"],
      "pcie": ["chip", "block", "PCIE_BK_csr"],
  }

  _REGISTERS = [
      "PARAM",
      "DMUXSEL",
      "txdata",
      "rxgate",
      "pad_info_PRESERVE",
      "rxdata",
      "rxdata_ISR",
      "rxdata_ISR_OVF",
      "rxdata_IER",
      "rxdata_IMR",
      "rxdata_ITR",
      "fiq",
  ]

  def __init__(self, name):
    self.name = "_".join(name.split("_")[:-1]).lower()
    self.pins = []
    self.functions = {}
    self.sswrp_data_name = "google_mbu_pinctrl_" + self.name
    self.pins_array_name = "google_mbu_" + self.name
    self.pin_groups_name = "google_mbu_" + self.name + "_groups"

    self.pin_functions_name = "google_mbu_" + self.name + "_functions"
    self.functions_enum_name = "google_mbu_pinmux_" + self.name + "_functions"

    self._si_codex = {}
    self.load_si_codex()

  def load_si_codex(self):
    with open(f"si_codex/{self._SI_CODEX_FILENAME[self.name]}", "r") as f:
      self._si_codex = json.load(f)
      for p in self._SI_CODEX_PATH[self.name]:
        self._si_codex = custom_container_get(self._si_codex, p)

  def pinctrl_pins_arr_str(self):
    ret = [
        f"static const struct pinctrl_pin_desc {self.pins_array_name}[] = {{"
    ]

    for i in range(0, len(self.pins)):
      ret.append(f'\tPINCTRL_PIN({i}, "{self.pins[i].pad_name}"),')

    ret.extend([
        "};",
        "",
    ])

    return ret

  def pin_groups_arr_str(self):
    ret = [f"static const struct google_pingroup {self.pin_groups_name}[] = {{"]
    wrapper = TextWrapper(
        width=_MAX_LINE_LENGTH, initial_indent="\t", subsequent_indent="\t\t"
    )

    for i in range(0, len(self.pins)):
      line = [f'PIN_GROUP({i}, "{self.pins[i].pad_name}"']

      for j in range(0, len(self.pins[i].pin_functions)):
        line.append(f", {self.pins[i].pin_functions[j]}")

      line.append("),")

      ret.extend(wrapper.wrap("".join(line)))

    ret.append("};\n")

    return ret

  def functions_enum_arr_str(self):
    ret = [f"enum {self.functions_enum_name} {{"]

    for func in self.functions:
      ret.append(f"\tgoogle_pinmux_{func},")

    ret.append("};\n")

    return ret

  def function_groups_arr_str(self):
    wrapper = TextWrapper(width=_MAX_LINE_LENGTH, subsequent_indent="\t\t")
    ret = []

    for func_name in self.functions:
      line = [f"FUNCTION_GROUPS({func_name}"]
      for pin in self.functions[func_name]:
        line.append(f', "{pin}"')

      line.append(");")

      ret.extend(wrapper.wrap("".join(line)))

    ret.append("")

    return ret

  def functions_arr_str(self):
    ret = [
        "static const struct google_pin_function "
        f"{self.pin_functions_name}[] = {{"
    ]

    for func in self.functions:
      ret.append(f"\tFUNCTION({func}),")

    ret.append("};")

    return ret

  def max_gpios_str(self):
    return f"\n#define MAX_NR_GPIO_{self.name.upper()} {len(self.pins)}\n"

  def sswrp_info_str(self):
    info = f"""\
      static const struct google_pinctrl_soc_sswrp_info {self.sswrp_data_name} = {{
      \t.pins = {self.pins_array_name},
      \t.num_pins = ARRAY_SIZE({self.pins_array_name}),
      \t.groups = {self.pin_groups_name},
      \t.num_groups = ARRAY_SIZE({self.pin_groups_name}),
      \t.funcs = {self.pin_functions_name},
      \t.num_funcs = ARRAY_SIZE({self.pin_functions_name}),
      \t.gpio_func = GPIO_FUNC_BIT_POS,
      \t.num_gpios = MAX_NR_GPIO_{self.name.upper()},
      \t.label = "{self.name}",
      \t.common = &mbu_common,
      #if IS_ENABLED(CONFIG_DEBUG_FS)
      \t.pins_excl_regs = {self.name}_pins_excl_regs,
      \t.npins_excl_regs = ARRAY_SIZE({self.name}_pins_excl_regs),
      #endif
      }};

    """

    return textwrap.dedent(info)

  def missing_regs(self):
    ret = collections.defaultdict(list)

    for pin in self.pins:
      for reg in self._REGISTERS:
        pin_reg_name = f"{pin.pad_name}_{reg}"
        pin_reg = custom_container_get(self._si_codex["reg"], pin_reg_name)

        if pin_reg is None:
          ret[pin.pad_name].append(reg)

    return ret

  def excluded_regs(self):
    map_registers = {
        "PARAM": "param",
        "DMUXSEL": "dmux",
        "txdata": "txdata",
        "rxgate": "rxgate",
        "rxdata": "rxdata",
        "rxdata_ISR": "isr",
        "rxdata_ISR_OVF": "isrovf",
        "rxdata_IER": "ier",
        "rxdata_IMR": "imr",
        "rxdata_ITR": "itr",
        "pad_info_PRESERVE": "pip",
        "fiq": "fiqts",
    }

    missing_regs = self.missing_regs()
    block = []

    for pin in missing_regs:
      l = []

      for reg in missing_regs[pin]:
        l.append(f"\t\t\t.{map_registers[reg]} = 1,")

      missing_pin = [
          "\t{",
          f"\t\t.pin_id = {pin},",
          "\t\t.layout = {",
          *l,
          "\t\t},",
          "\t},",
      ]

      block.extend(missing_pin)

    if not block:
      return block

    res = [
        "#if IS_ENABLED(CONFIG_DEBUG_FS)",
        (
            f"static const struct google_pinctrl_registers_flags"
            f" {self.name}_pins_excl_regs[] = {{"
        ),
        *block,
        "};",
        "#endif",
    ]

    return res


class Common:
  """Contains utility functions that are not specific to a particulat SSWRP."""

  def __init__(self, sswrp_list):
    self.sswrp_list = sswrp_list

  # pylint: disable=missing-function-docstring
  def compatible_arr_str(self):

    def compatible_single_arr_str(sswrp):
      compatible_string = "google," + "mbu-" + sswrp.name + "-pinctrl"
      compatible_name = compatible_string.replace("_", "-")

      ret = [
          f'\t{{ .compatible = "{compatible_name}",',
          f"\t  .data = &{sswrp.sswrp_data_name} }},",
      ]

      return ret

    match_table_name = "google_mbu_of_match"

    ret = [f"static const struct of_device_id {match_table_name}[] = {{"]
    for sswrp in self.sswrp_list:
      ret.extend(compatible_single_arr_str(sswrp))

    # Add sentinel
    ret.append("\t{},")

    ret.append("};")

    return ret

  def google_pins_arr_str(self):
    max_pins = 0
    for sswrp in self.sswrp_list:
      max_pins = max(max_pins, len(sswrp.pins))

    ret = []
    for i in range(0, max_pins):
      ret.append(f"GOOGLE_PINS({i});")

    ret.append("")

    return ret

  def header_str(self):
    header = """
      // SPDX-License-Identifier: GPL-2.0-only
      /*
       * Copyright 2023 Google LLC
       *
       * Please do not edit this file.
       * It was generated using mbu-parser.py.
       */
      #include <linux/module.h>
      #include <linux/of.h>
      #include <linux/pinctrl/pinctrl.h>
      #include <linux/platform_device.h>
      #include <dt-bindings/gpio/google,mbu.h>

      #include "common.h"
      #define google_pinmux__ -1

      static const u8 mbu_regs[] = {
      \tPARAM_ID,
      \tDMUX_ID,
      \tTXDATA_ID,
      \tRXGATE_ID,
      \tRXDATA_ID,
      \tISR_ID,
      \tISROVF_ID,
      \tIER_ID,
      \tIMR_ID,
      \tITR_ID,
      \tPIP_ID,
      \tFIQTS_ID,
      };

      static const u8 mbu_reg2offset[REGS_NUM] = {
      \t[PARAM_ID] =  0x00,
      \t[DMUX_ID] =   0x04,
      \t[TXDATA_ID] = 0x08,
      \t[RXGATE_ID] = 0x0C,
      \t[RXDATA_ID] = 0x1C,
      \t[ISR_ID] =    0x20,
      \t[ISROVF_ID] = 0x24,
      \t[IER_ID] =    0x28,
      \t[IMR_ID] =    0x2C,
      \t[ITR_ID] =    0x30,
      \t[PIP_ID] =    0x40,
      \t[FIQTS_ID] =  0x44,
      };

      #define MBU_DUMP_REGS_HEADER " idx |          SOC pin           |   PARAM  |   DMUX   |  TXDATA  |  RXGATE  |  RXDATA  |   ISR    |  ISROVF  |   IER    |   IMR    |    ITR   |PADINFOPRE|   FIQTS  |\\n"
      static_assert(sizeof(MBU_DUMP_REGS_HEADER) <= DUMP_REGS_BUF_SIZE);

      #define MBU_PIN_SIZE 0x1000

      static const struct platform_common_info mbu_common = {
      \t.regs = mbu_regs,
      \t.num_regs = ARRAY_SIZE(mbu_regs),
      \t.reg2offset = mbu_reg2offset,
      \t.pin_stride = MBU_PIN_SIZE,
      #if IS_ENABLED(CONFIG_DEBUG_FS)
      \t.dump_regs_header = MBU_DUMP_REGS_HEADER,
      #endif
      };
    """

    return textwrap.dedent(header[1:])

  def footer_str(self):
    footer = """
      static int google_mbu_pinctrl_probe(struct platform_device *pdev)
      {
      \treturn google_pinctrl_probe(pdev, google_mbu_of_match);
      }

      static struct platform_driver google_mbu_pinctrl_driver = {
      \t.driver = {
      \t\t.name = "google_mbu_pinctrl",
      \t\t.of_match_table = google_mbu_of_match,
      #ifdef CONFIG_PM
      \t\t.pm = &google_pinctrl_pm_ops,
      #endif
      \t},
      \t.probe = google_mbu_pinctrl_probe,
      \t.remove = google_pinctrl_remove,
      };

      static int __init google_mbu_pinctrl_init(void)
      {
      \treturn platform_driver_register(&google_mbu_pinctrl_driver);
      }
      arch_initcall(google_mbu_pinctrl_init);

      static void __exit google_mbu_pinctrl_exit(void)
      {
      \tplatform_driver_unregister(&google_mbu_pinctrl_driver);
      }
      module_exit(google_mbu_pinctrl_exit);

      MODULE_AUTHOR("Google LLC");
      MODULE_DESCRIPTION("Google MBU PINCTRL Driver");
      MODULE_LICENSE("GPL");
    """

    return textwrap.dedent(footer)

  def pinctrl_arr_str(self):
    ret = [
        self.header_str(),
        *self.google_pins_arr_str(),
    ]

    for sswrp in self.sswrp_list:
      ret.extend([
          *sswrp.pinctrl_pins_arr_str(),
          *sswrp.functions_enum_arr_str(),
          *sswrp.function_groups_arr_str(),
          *sswrp.pin_groups_arr_str(),
          *sswrp.functions_arr_str(),
          *sswrp.excluded_regs(),
          sswrp.max_gpios_str(),
          sswrp.sswrp_info_str(),
      ])

    ret.extend([
        *self.compatible_arr_str(),
        self.footer_str(),
    ])

    return ret

  def print_pinctrl(self):
    print("\n".join(self.pinctrl_arr_str()))


class Parser:
  """Handles the parsing from IO sheet."""

  def __init__(self):
    self.wb = load_workbook("mbu_data_sheet.xlsx")
    self.ws = self.wb["IO pads"]
    self.row_count = self.ws.max_row
    # The SSWRPs for which we want to extract the pin information
    self.valid_sswrps = [
        "HSIOS_BK",
        "HSIOS_STBY_BK",
        "LSIOS_BK",
        "LSIOE_BK",
        "PCIE_BK",
        "AOSS_BK",
        "GDMC_BK",
        "AOSS_AUDIO_BK",
    ]
    # At least one pin of this SSWRP has been found
    self.accessed_sswrps = []
    self.sswrp_list = []
    # List of valid functions for every SSWRP
    self.valid_functions = {
        "hsios": [
            "gpio", "ufs", "sd_data", "sd_cmd", "sd_fbclk", "sd_clk",
            "atb1",
        ],
        "hsios_stby": [
            "clkbuf", "ufs"
        ],
        "lsios": [
            "gpio", "mclk", "debug_mux", "qspi0", "pwm",
            "dpu", "vsync", "camera_mute", "pre_ocp_gpu", "soft_pre_ocp_gpu",
            "spi0", "spi1", "spi2", "spi3", "i3c0", "i3c1", "i3c2", "i3c3",
            "i3c4", "i2c0", "i2c1", "i2c2", "i2c3", "i2c4", "uart0",
            "uart1","uart2", "uart3"
        ],
        "lsioe": [
            "gpio", "i2c", "i3c", "spi4", "uart4", "uart5", "spi5"
        ],
        "pcie": [
            "pcie0", "atb0", "gpio"
        ],
        "gdmc": [
            "gpio", "uart", "debug", "jtag", "cti", "tm0", "tm1"
        ],
        "aoss": [
            "gpio", "uart0", "uart1", "uart2", "uart3", "uart4", "uart5", "uart6",
            "uart7", "spi0", "spi1", "spi2", "spi3", "spi4" , "spi5", "spi6",
            "spi7", "i2c0", "i2c1", "i2c2", "i2c3", "i2c4", "i2c5","i3c0",
            "i3c1", "i3c2", "i3c3", "i3c4", "i3c5", "i3c6", "i3c7", "cpm"
        ],
        "aoss_audio": [
            "gpio", "pdm0", "pdm1", "pdm2", "pdm3", "i2s0", "i2s1", "tdm0",
            "tdm1", "sdwire0", "sdwire1", "sdwire2", "xtal", "out", "stby",
        ]
    }

  def get_short_fname(self, long_fname, sswrp_name, function_attrs):
    """Eliminates the pin information from the function names.

    The function names in the IO sheet contain pin information in them, this
    function removes it.
    eg. Converts lsios_cli3_pin0 to lsios_cli3.

    Args:
        long_fname: The long function name taken from the IO sheet, eg.
          lsios_cli3_pin0.
        sswrp_name: The name of the sswrp to which this pin belongs.
        function_attrs: The attributes associated with the function, which may
          contain information like "GPIO=1".

    Returns:
        Short name of the function, eg. lsios_cli3.
    """
    if "GPIO=1" in function_attrs:
      return f"{sswrp_name}_gpio"

    short_fnames = []
    for func in self.valid_functions[sswrp_name]:
      if func in long_fname:
        short_fnames.append(func)
    if not short_fnames:
      sys.exit(
          "No matching short function name present for {}".format(long_fname)
      )
    elif len(short_fnames) > 1:
      # Choosing the longest matching short name
      short_fnames.sort(key=lambda x: (-len(x), x))

    return sswrp_name + "_" + short_fnames[0]

  def get_functions(self, row, pad_name):
    """Parses the functions for the pin on the nth row.

    Parses and stores the supported functions for the pin on the
    row'th row. It also creates an alternate mapping between the
    function name and the pins that can support it.

    Args:
        row: The query row in the IO sheet.
        pad_name: String that contains the name of the pin on the row'th row.

    Returns:
        An array that contains all the functions supported
        by the pin on the given row.
    """
    pin_functions = []
    func0_col = 40
    max_nr_funcs = 9

    for i in range(0, max_nr_funcs):
      col_num = func0_col + i * 2
      function_cell = get_column_letter(col_num) + str(row)
      function_attrs_cell = get_column_letter(col_num + 1) + str(row)
      function_attrs = self.ws[function_attrs_cell].value or ""

      if self.ws[function_cell].value is not None and str(self.ws[function_cell].value.lower()) != "no" :
        long_fname = self.ws[function_cell].value.lower()
        short_fname = self.get_short_fname(
            long_fname, self.sswrp_list[-1].name, function_attrs
        )
        pin_functions.append(short_fname)

        # For every function, populating the pins that support it
        if short_fname not in self.sswrp_list[-1].functions:
          self.sswrp_list[-1].functions[short_fname] = []
        self.sswrp_list[-1].functions[short_fname].append(pad_name)
      else:
        pin_functions.append("_")
    return pin_functions

  # Gets executed when we access a valid SSWRP for the first time
  def check_and_append_sswrp(self, sswrp_name):
    if sswrp_name not in self.accessed_sswrps:
      self.sswrp_list.append(Sswrp(sswrp_name))
      self.accessed_sswrps.append(sswrp_name)

  # A valid pin belongs to a valid SSWRP and has a non empty pad number field
  def is_pin_valid(self, pad_nr_cell, sswrp_name):
    return (sswrp_name in self.valid_sswrps) and (self.ws[pad_nr_cell].value
                                                  is not None)

  def check_if_every_function_used(self):
    """Verify if all valid functions are utilized within each SSWRP.

    This check helps ensure the completeness and accuracy of the parsed
    data, highlighting any potential omissions in the pin function
    mappings.
    """
    for sswrp in self.sswrp_list:
      funcs = list(sswrp.functions.keys())
      valid_funcs = [
          f"{sswrp.name}_{x}" for x in self.valid_functions[sswrp.name]
      ]

      diff = set(valid_funcs) - set(funcs)
      if diff:
        print(
            f"Some functions have not been used in {sswrp.name}: {diff}",
            file=sys.stderr,
        )

  def check_if_every_sswrps_used(self):
    """Verify if all valid SSWRPs are present in the parsed data.

    This check helps ensure that all intended SSWRPs are included in the
    generated pinctrl driver code, identifying any potential omissions
    in the parsing process.
    """
    sswrps = [x.name for x in self.sswrp_list]
    valid_sswrps = [x.replace("_BK", "").lower() for x in self.valid_sswrps]
    diff = set(valid_sswrps) ^ set(sswrps)
    if diff:
      print(f"Some SSWRP have not been used: {diff}", file=sys.stderr)

    diff = set(self.valid_functions.keys()) ^ set(valid_sswrps)
    if diff:
      print(
          "Some SSWRP have defined functions but are not used completely:"
          f" {diff}",
          file=sys.stderr,
      )

  def parse_data(self):
    """The main function used to parse data."""
    for row in range(1, self.row_count + 1):
      sswrp_cell = "K" + str(row)
      sswrp_name = self.ws[sswrp_cell].value
      pad_nr_cell = "C" + str(row)
      pad_name_cell = "D" + str(row)
      pad_name = self.ws[pad_name_cell].value

      if self.is_pin_valid(pad_nr_cell, sswrp_name):
        self.check_and_append_sswrp(sswrp_name)
        self.sswrp_list[-1].pins.append(
            Pin(self.ws[pad_nr_cell].value, pad_name,
                self.get_functions(row, pad_name)))

    self.check_if_every_function_used()
    self.check_if_every_sswrps_used()

    return self.sswrp_list


def main():

  arg_parser = argparse.ArgumentParser()
  arg_parser.add_argument("-f", "--out_file", help="Writes output to file")
  args = arg_parser.parse_args()
  if args.out_file:
    sys.stdout = open(args.out_file, "w")

  parser = Parser()
  sswrp_list = parser.parse_data()

  common = Common(sswrp_list)
  common.print_pinctrl()


if __name__ == "__main__":
  main()
